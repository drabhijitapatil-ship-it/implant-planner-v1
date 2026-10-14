"""Research Module — pure engines (no DB access).

* Case Field Registry  — the only case fields a study may link / query.
* Condition grammar    — eligibility rules + show_if  → pass / fail / unknown.
* Formula evaluator    — safe derived-variable expressions (no eval).
* Schedule engine      — timepoint → due date / window / status.
* Descriptive stats    — per variable × timepoint summaries.
* CBL study template   — the seed fixture adapted to the real Implanr schema.

Studies live in db.research_studies, enrolments in db.research_enrolments
(forms embedded), patient-level values in db.research_patient_values.
"""
from __future__ import annotations

import io
import math
import re
import statistics
from datetime import date, datetime, timedelta, timezone
from typing import Any, Callable, Optional

STUDY_TYPES = ["prospective_cohort", "retrospective_cohort", "rct", "case_series", "split_mouth", "cross_sectional"]
STUDY_STATUSES = ["draft", "active", "recruitment_closed", "follow_up_complete", "data_locked", "archived"]
LEVELS = ["patient", "implant", "site"]
SOURCES = ["linked", "study", "derived"]
DATA_TYPES = ["integer", "decimal", "text", "single_select", "multi_select", "boolean", "date", "scale"]
ANCHORS = ["enrolment", "placement_date", "second_stage_date", "prosthesis_delivery_date", "variable"]
MEMBER_ROLES = ["pi", "guide", "co_investigator", "data_collector", "statistician"]
ENROLMENT_STATUSES = ["enrolled", "active", "completed", "withdrawn", "lost_to_follow_up", "excluded_post_enrolment"]
CODE_RE = re.compile(r"^[A-Z0-9-]{3,20}$")
KEY_RE = re.compile(r"^[a-z][a-z0-9_]{0,39}$")
MISSED_MULTIPLIER = 2  # BR-08 default: > 2× window_after past due → missed


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def today() -> date:
    return datetime.now(timezone.utc).date()


# --------------------------------------------------------------------------
# Case Field Registry  (path → how to read it from a procedure / implant)
# --------------------------------------------------------------------------
YES_NO = [{"code": "Yes", "label": "Yes"}, {"code": "No", "label": "No"}]
_PROC_TYPES = ["Single Conventional Implant", "Multiple Conventional Implants", "Immediate Implant", "Sinus Lift",
               "Implant Overdenture", "All on 4", "All on 6", "All on X",
               "Implant Placement with Guided Bone Regeneration", "Existing Implant", "Zygomatic / Pterygoid"]


def _opts(vals):
    return [{"code": v, "label": v} for v in vals]


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    m = re.search(r"-?\d+(\.\d+)?", str(v))
    return float(m.group(0)) if m else None


def _str(v):
    return None if v in (None, "") else str(v)


def _list(v):
    return v if isinstance(v, list) else ([v] if v not in (None, "") else None)


def _tooth_key(implant):
    return str((implant or {}).get("tooth_number") or "")


def _pd(proc, implant, k):
    bpd = proc.get("baseline_probing_depths") or {}
    row = bpd.get(_tooth_key(implant)) or {}
    return _num(row.get(k))


def _last_followup(proc):
    fus = [f for f in (proc.get("followups") or []) if isinstance(f, dict)]
    return fus[-1] if fus else {}


def _survival(proc, implant):
    fu = _last_followup(proc)
    rev = (fu.get("survival_review") or {}).get(_tooth_key(implant)) or {}
    return _str(rev.get("status"))


def _fu_pd(proc, implant, k):
    fu = _last_followup(proc)
    row = (fu.get("probing_depths") or {}).get(_tooth_key(implant)) or {}
    return _num(row.get(k))


def _bop(proc, implant):
    fu = _last_followup(proc)
    st = ((fu.get("soft_tissue") or {}).get("bleeding_on_probing") or {}).get("status")
    return None if st is None else st


def _med(k):
    return lambda p, i: _str((p.get("medical_assessment") or {}).get(k))


def _p(k, conv=_str):
    return lambda p, i: conv(p.get(k))


def _i(k, conv=_str):
    return lambda p, i: conv((i or {}).get(k))


# (path, label, level, data_type, unit, options, phase, getter)
_REG: list[tuple] = [
    ("patient.age_years", "Patient age", "patient", "integer", "years", None, 1, _p("age", _num)),
    ("patient.sex", "Sex", "patient", "single_select", None, _opts(["Male", "Female", "Other"]), 1, _p("sex")),
    ("patient.medical.diabetes", "Diabetes", "patient", "single_select", None, _opts(["No", "Controlled", "Uncontrolled"]), 1, _med("diabetes")),
    ("patient.medical.smoking", "Smoking status", "patient", "single_select", None, _opts(["No", "Light (<10/day)", "Heavy (>10/day)"]), 1, _med("smoking")),
    ("patient.medical.anticoagulant", "On anticoagulants", "patient", "single_select", None, YES_NO, 1, _med("anticoagulant")),
    ("patient.medical.osteoporosis", "Osteoporosis / bisphosphonates", "patient", "single_select", None, YES_NO, 1, _med("osteoporosis")),
    ("patient.medical.radiation", "Head & neck radiation", "patient", "single_select", None, YES_NO, 1, _med("radiation")),
    ("patient.medical.hba1c_pct", "HbA1c", "patient", "decimal", "%", None, 1, lambda p, i: _num((p.get("medical_assessment") or {}).get("hba1c"))),
    ("patient.medical.cigarettes_per_day", "Cigarettes per day", "patient", "integer", "/day", None, 1,
     lambda p, i: _num((p.get("medical_assessment") or {}).get("cigarettes_per_day"))),
    ("patient.medical.bisphosphonate_history", "Bisphosphonate history", "patient", "single_select", None,
     _opts(["No", "Oral", "IV"]), 1, _med("bisphosphonate")),
    ("patient.medical_risk_level", "Medical risk level", "patient", "single_select", None,
     _opts(["Low Risk", "Moderate Risk", "High Risk"]), 1, _p("medical_risk_level")),
    ("case.procedure_type", "Procedure type", "patient", "single_select", None, _opts(_PROC_TYPES), 1, _p("implant_procedure_type")),
    ("case.arch", "Arch", "patient", "single_select", None, _opts(["Maxilla", "Mandible", "Both"]), 1, _p("arch")),
    ("case.status", "Case status", "patient", "text", None, None, 0, _p("status")),
    ("case.procedure_date", "Procedure (planned) date", "patient", "date", None, None, 1, _p("procedure_date")),
    ("case.periodontal_status", "Periodontal status", "patient", "text", None, None, 1, _p("periodontal_status")),
    ("case.gingival_biotype", "Gingival biotype", "patient", "single_select", None, _opts(["Thin", "Thick"]), 1, _p("gingival_biotype")),
    ("case.smile_line", "Smile line", "patient", "single_select", None, _opts(["Low", "Medium", "High"]), 1, _p("smile_line")),
    ("case.keratinized_mucosa", "Keratinized mucosa", "patient", "text", None, None, 1, _p("keratinized_mucosa")),
    ("case.parafunction_habit", "Parafunctional habit", "patient", "text", None, None, 1, _p("parafunction_habit")),
    ("case.aesthetic_risk.overall", "Aesthetic risk (ERA) overall", "patient", "single_select", None,
     _opts(["Low", "Medium", "High"]), 1, lambda p, i: _str((p.get("aesthetic_risk") or {}).get("overall_risk"))),
    ("case.loading_type", "Loading protocol", "patient", "multi_select", None, None, 1, _p("loading_type", _list)),
    ("case.augmentation_required", "Augmentation required", "patient", "boolean", None, None, 1, _p("augmentation_required", lambda v: v if isinstance(v, bool) else None)),
    ("case.sinus_lift_type", "Sinus lift type", "patient", "text", None, None, 1, _p("sinus_lift_type")),
    ("case.prosthetic_plan", "Prosthetic plan", "patient", "text", None, None, 1, _p("prosthetic_plan")),
    ("case.missing_teeth", "Missing teeth (FDI)", "patient", "multi_select", None, None, 1, _p("missing_teeth", _list)),
    ("phase2.actual_surgery_date", "Surgery date (Phase 2)", "patient", "date", None, None, 2, _p("phase2_actual_done_date")),
    ("phase3.isq", "ISQ at second stage", "patient", "integer", "ISQ", None, 3, lambda p, i: _num((p.get("phase3_data") or {}).get("isq_value"))),
    ("phase3.healing_abutment_height", "Healing abutment height", "patient", "decimal", "mm", None, 3,
     lambda p, i: _num((p.get("phase3_data") or {}).get("healing_abutment_height"))),
    ("phase3.done_date", "Second-stage date", "patient", "date", None, None, 3, _p("phase3_done_date")),
    ("phase4.prosthesis_delivery_date", "Prosthesis delivery date", "patient", "date", None, None, 4, _p("phase4_step2_done_date")),
    ("phase4.final_prosthetic_plan", "Final prosthesis", "patient", "text", None, None, 4, _p("final_prosthetic_plan")),
    # implant level
    ("implant.site_fdi", "Implant site (FDI)", "implant", "integer", None, None, 2, _i("tooth_number", _num)),
    ("implant.brand", "Implant brand", "implant", "text", None, None, 2, _i("brand")),
    ("implant.system", "Implant system", "implant", "text", None, None, 2, _i("system")),
    ("implant.diameter_mm", "Implant diameter", "implant", "decimal", "mm", None, 2, _i("diameter", _num)),
    ("implant.length_mm", "Implant length", "implant", "decimal", "mm", None, 2, _i("length", _num)),
    ("implant.bone_type", "Bone type (Lekholm & Zarb)", "implant", "single_select", None, _opts(["D1", "D2", "D3", "D4"]), 2, _i("bone_type")),
    ("implant.bone_width_mm", "Bone width", "implant", "decimal", "mm", None, 2, _i("bone_width", _num)),
    ("implant.bone_height_mm", "Bone height", "implant", "decimal", "mm", None, 2, _i("bone_height", _num)),
    ("implant.insertion_torque_ncm", "Insertion torque", "implant", "integer", "Ncm", None, 2, _i("insertion_torque_ncm", _num)),
    ("implant.placement_date", "Placement date", "implant", "date", None, None, 2, _i("placement_date")),
    # site level
    ("site.socket_type", "Extraction socket type (Elian)", "site", "single_select", None, _opts(["Type I", "Type II", "Type III"]), 1,
     lambda p, i: _str((p.get("socket_types") or {}).get(_tooth_key(i)))),
    ("site.baseline_pd_mesial", "Baseline PD mesial", "site", "decimal", "mm", None, 4, lambda p, i: _pd(p, i, "mesial")),
    ("site.baseline_pd_distal", "Baseline PD distal", "site", "decimal", "mm", None, 4, lambda p, i: _pd(p, i, "distal")),
    ("site.baseline_pd_vestibular", "Baseline PD vestibular", "site", "decimal", "mm", None, 4, lambda p, i: _pd(p, i, "vestibular")),
    ("site.baseline_pd_lingual", "Baseline PD lingual", "site", "decimal", "mm", None, 4, lambda p, i: _pd(p, i, "lingual")),
    ("site.baseline_kgw", "Baseline keratinised gingiva width", "site", "decimal", "mm", None, 4, lambda p, i: _pd(p, i, "kgw")),
    ("site.latest_pd_mesial", "Latest follow-up PD mesial", "site", "decimal", "mm", None, 5, lambda p, i: _fu_pd(p, i, "mesial")),
    ("site.latest_pd_distal", "Latest follow-up PD distal", "site", "decimal", "mm", None, 5, lambda p, i: _fu_pd(p, i, "distal")),
    ("site.latest_bop", "Bleeding on probing (latest follow-up)", "site", "single_select", None, _opts(["Present", "Absent"]), 5, _bop),
    ("site.survival_status", "Implant survival status (latest review)", "site", "single_select", None,
     _opts(["Surviving", "Failed", "Removed"]), 5, _survival),
]
REGISTRY: dict[str, dict] = {}
for _path, _label, _level, _dt, _unit, _options, _phase, _getter in _REG:
    REGISTRY[_path] = {"path": _path, "label": _label, "level": _level, "data_type": _dt, "unit": _unit,
                       "options": _options, "phase": _phase, "_get": _getter}


def registry_public() -> list[dict]:
    return [{k: v for k, v in r.items() if k != "_get"} for r in REGISTRY.values()]


def case_value(proc: dict, implant: Optional[dict], path: str):
    reg = REGISTRY.get(path)
    if not reg:
        return None
    try:
        return reg["_get"](proc or {}, implant or {})
    except Exception:
        return None


def case_sites(proc: dict) -> list[dict]:
    """Implant/site candidates of a case: placed implants first, else planned
    missing teeth (implant data unknown)."""
    sites = []
    for idx, imp in enumerate(proc.get("implants") or []):
        if isinstance(imp, dict) and imp.get("tooth_number") not in (None, ""):
            sites.append({"index": idx, "tooth_number": str(imp["tooth_number"]), "implant": imp, "planned": False})
    if sites:
        return sites
    for idx, t in enumerate(proc.get("missing_teeth") or []):
        plan = next((pl for pl in (proc.get("implant_plans") or []) if str(pl.get("position")) == str(t)), None)
        imp = {"tooth_number": str(t)}
        if plan:
            imp.update({"brand": plan.get("brand"), "system": plan.get("system"), "diameter": plan.get("diameter"),
                        "length": plan.get("length"), "bone_type": plan.get("bone_type"),
                        "bone_width": plan.get("bone_width"), "bone_height": plan.get("bone_height")})
        sites.append({"index": idx, "tooth_number": str(t), "implant": imp, "planned": True})
    return sites


def anchor_date(proc: dict, implant: Optional[dict], anchor: str, enrolled_at: Optional[str] = None,
                var_value: Any = None) -> Optional[date]:
    if anchor == "enrolment":
        return _to_date(enrolled_at)
    if anchor == "placement_date":
        return _to_date((implant or {}).get("placement_date")) or _to_date(proc.get("phase2_actual_done_date"))
    if anchor == "second_stage_date":
        return _to_date(proc.get("phase3_done_date"))
    if anchor == "prosthesis_delivery_date":
        return _to_date(proc.get("phase4_step2_done_date"))
    if anchor == "variable":
        return _to_date(var_value)
    return None


def _to_date(v) -> Optional[date]:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


# --------------------------------------------------------------------------
# Condition grammar  → pass / fail / unknown
# --------------------------------------------------------------------------
def _cmp_val(v):
    n = _num(v)
    return n if n is not None else (str(v).strip().lower() if v is not None else None)


def evaluate_condition(cond: Any, getter: Callable[[str], Any]) -> str:
    if not isinstance(cond, dict):
        return "unknown"
    if "all" in cond:
        res = [evaluate_condition(c, getter) for c in (cond["all"] or [])]
        if not res:
            return "unknown"
        if "fail" in res:
            return "fail"
        return "unknown" if "unknown" in res else "pass"
    if "any" in cond:
        res = [evaluate_condition(c, getter) for c in (cond["any"] or [])]
        if not res:
            return "unknown"
        if "pass" in res:
            return "pass"
        return "unknown" if "unknown" in res else "fail"
    if "not" in cond:
        r = evaluate_condition(cond["not"], getter)
        return {"pass": "fail", "fail": "pass"}.get(r, "unknown")
    field, op, target = cond.get("field"), cond.get("op"), cond.get("value")
    if not field or not op:
        return "unknown"
    actual = getter(field)
    if op == "exists":
        return "pass" if actual not in (None, "", []) else "fail"
    if op == "not_exists":
        return "fail" if actual not in (None, "", []) else "pass"
    if actual in (None, "", []):
        return "unknown"
    try:
        if op in ("in", "not_in"):
            targets = {_cmp_val(t) for t in (target if isinstance(target, list) else [target])}
            actuals = [_cmp_val(a) for a in (actual if isinstance(actual, list) else [actual])]
            hit = any(a in targets for a in actuals)
            return "pass" if (hit if op == "in" else not hit) else "fail"
        if op == "contains":
            hay = [str(a).lower() for a in (actual if isinstance(actual, list) else [actual])]
            return "pass" if any(str(target).lower() in h for h in hay) else "fail"
        a, t = _cmp_val(actual if not isinstance(actual, list) else actual[0]), _cmp_val(target)
        if op == "eq":
            return "pass" if a == t else "fail"
        if op == "neq":
            return "pass" if a != t else "fail"
        if op == "between":
            lo, hi = _num(target[0]), _num(target[1])
            return "pass" if isinstance(a, (int, float)) and lo <= a <= hi else "fail"
        if not isinstance(a, (int, float)) or not isinstance(t, (int, float)):
            return "unknown"
        return "pass" if {"gt": a > t, "gte": a >= t, "lt": a < t, "lte": a <= t}[op] else "fail"
    except Exception:
        return "unknown"


def evaluate_eligibility(study: dict, proc: dict, implant: Optional[dict]) -> dict:
    """Per-rule results + overall: eligible | excluded | needs_review."""
    getter = lambda path: case_value(proc, implant, path)  # noqa: E731
    rules_out, overall = [], "eligible"
    for idx, rule in enumerate(study.get("eligibility") or []):
        kind, mode = rule.get("kind", "include"), rule.get("mode", "query")
        if mode == "checklist":
            rules_out.append({"index": idx, "kind": kind, "mode": mode, "text": rule.get("checklist_text", ""), "result": "checklist"})
            continue
        r = evaluate_condition(rule.get("query"), getter)
        rules_out.append({"index": idx, "kind": kind, "mode": mode, "result": r, "label": describe_condition(rule.get("query"))})
        if kind == "include" and r == "fail":
            overall = "excluded"
        elif kind == "exclude" and r == "pass":
            overall = "excluded"
        elif r == "unknown" and overall != "excluded":
            overall = "needs_review"
    return {"rules": rules_out, "overall": overall}


_OP_TXT = {"eq": "=", "neq": "≠", "in": "in", "not_in": "not in", "gt": ">", "gte": "≥", "lt": "<", "lte": "≤",
           "between": "between", "exists": "is recorded", "not_exists": "is not recorded", "contains": "contains"}


def describe_condition(cond: Any) -> str:
    if not isinstance(cond, dict):
        return ""
    for grp, joiner in (("all", " AND "), ("any", " OR ")):
        if grp in cond:
            return "(" + joiner.join(describe_condition(c) for c in cond[grp]) + ")"
    if "not" in cond:
        return "NOT " + describe_condition(cond["not"])
    label = (REGISTRY.get(cond.get("field") or "") or {}).get("label") or cond.get("field", "?")
    op = _OP_TXT.get(cond.get("op"), cond.get("op", ""))
    v = cond.get("value")
    vs = ", ".join(str(x) for x in v) if isinstance(v, list) else ("" if v is None else str(v))
    return f"{label} {op} {vs}".strip()


def validate_condition(cond: Any) -> list[str]:
    errs = []
    if not isinstance(cond, dict):
        return ["Condition must be an object"]
    for grp in ("all", "any"):
        if grp in cond:
            if not isinstance(cond[grp], list) or not cond[grp]:
                errs.append(f"'{grp}' needs at least one condition")
            else:
                for c in cond[grp]:
                    errs += validate_condition(c)
            return errs
    if "not" in cond:
        return validate_condition(cond["not"])
    if cond.get("field") not in REGISTRY:
        errs.append(f"Unknown case field '{cond.get('field')}'")
    if cond.get("op") not in _OP_TXT:
        errs.append(f"Unknown operator '{cond.get('op')}'")
    return errs


# --------------------------------------------------------------------------
# Formula evaluator (safe, recursive descent)
# --------------------------------------------------------------------------
_TOKEN = re.compile(r"\s*(?:(\d+\.?\d*)|('[^']*')|([A-Za-z_][A-Za-z0-9_]*)|(<=|>=|==|!=|[-+*/(),<>]))")


class Missing(Exception):
    pass


def _tokenize(src: str) -> list[tuple[str, Any]]:
    out, pos = [], 0
    src = src.strip()
    while pos < len(src):
        m = _TOKEN.match(src, pos)
        if not m or m.end() == pos:
            raise ValueError(f"Bad token near '{src[pos:pos+10]}'")
        pos = m.end()
        num, strv, ident, sym = m.groups()
        if num is not None:
            out.append(("num", float(num)))
        elif strv is not None:
            out.append(("str", strv[1:-1]))
        elif ident is not None:
            out.append(("id", ident))
        else:
            out.append(("sym", sym))
    return out


class _Parser:
    def __init__(self, tokens, lookup):
        self.t, self.i, self.lookup = tokens, 0, lookup

    def peek(self):
        return self.t[self.i] if self.i < len(self.t) else (None, None)

    def take(self, kind=None, val=None):
        k, v = self.peek()
        if kind and k != kind or (val is not None and v != val):
            raise ValueError(f"Expected {val or kind}, got {v}")
        self.i += 1
        return v

    def parse(self):
        v = self.expr_or()
        if self.i != len(self.t):
            raise ValueError("Unexpected trailing tokens")
        return v

    def expr_or(self):
        v = self.expr_and()
        while self.peek() == ("id", "or"):
            self.take(); r = self.expr_and(); v = bool(v) or bool(r)
        return v

    def expr_and(self):
        v = self.expr_not()
        while self.peek() == ("id", "and"):
            self.take(); r = self.expr_not(); v = bool(v) and bool(r)
        return v

    def expr_not(self):
        if self.peek() == ("id", "not"):
            self.take(); return not self.expr_not()
        return self.expr_cmp()

    def expr_cmp(self):
        v = self.expr_add()
        k, s = self.peek()
        if k == "sym" and s in ("<", ">", "<=", ">=", "==", "!="):
            self.take(); r = self.expr_add()
            return {"<": v < r, ">": v > r, "<=": v <= r, ">=": v >= r, "==": v == r, "!=": v != r}[s]
        return v

    def expr_add(self):
        v = self.expr_mul()
        while self.peek()[0] == "sym" and self.peek()[1] in ("+", "-"):
            s = self.take(); r = self.expr_mul(); v = v + r if s == "+" else v - r
        return v

    def expr_mul(self):
        v = self.expr_unary()
        while self.peek()[0] == "sym" and self.peek()[1] in ("*", "/"):
            s = self.take(); r = self.expr_unary()
            if s == "/":
                if r == 0:
                    raise Missing("division by zero")
                v = v / r
            else:
                v = v * r
        return v

    def expr_unary(self):
        if self.peek() == ("sym", "-"):
            self.take(); return -self.expr_unary()
        return self.atom()

    def atom(self):
        k, v = self.peek()
        if k == "num":
            self.take(); return v
        if k == "str":
            self.take(); return v
        if k == "sym" and v == "(":
            self.take(); r = self.expr_or(); self.take("sym", ")"); return r
        if k == "id":
            self.take()
            if v in ("true", "false"):
                return v == "true"
            if self.peek() == ("sym", "("):
                self.take()
                args = []
                if self.peek() != ("sym", ")"):
                    args.append(self.expr_or())
                    while self.peek() == ("sym", ","):
                        self.take(); args.append(self.expr_or())
                self.take("sym", ")")
                return self.call(v, args)
            raise ValueError(f"Unknown identifier '{v}'")
        raise ValueError("Unexpected end of expression")

    def call(self, name, args):
        if name == "v":
            val = self.lookup(str(args[0]), str(args[1]) if len(args) > 1 else None)
            if val in (None, ""):
                raise Missing(args[0])
            n = _num(val)
            return n if n is not None else val
        if name == "is_missing":
            # special: argument evaluated already; if it raised Missing we'd not get here
            return False
        if name == "mean":
            return sum(args) / len(args)
        if name == "abs":
            return abs(args[0])
        if name == "min":
            return min(args)
        if name == "max":
            return max(args)
        if name == "round":
            return round(args[0], int(args[1]) if len(args) > 1 else 0)
        if name == "if":
            return args[1] if args[0] else args[2]
        if name == "days_between":
            d1, d2 = _to_date(args[0]), _to_date(args[1])
            if not d1 or not d2:
                raise Missing("date")
            return (d2 - d1).days
        raise ValueError(f"Unknown function '{name}'")


def evaluate_formula(formula: str, lookup: Callable[[str, Optional[str]], Any]):
    """Returns (value, error). value is None when any input is missing."""
    try:
        val = _Parser(_tokenize(formula), lookup).parse()
        if isinstance(val, float):
            if math.isnan(val) or math.isinf(val):
                return None, None
            val = round(val, 4)
        return val, None
    except Missing:
        return None, None
    except Exception as e:  # syntax errors
        return None, str(e)


def validate_formula(formula: str, known_keys: set[str]) -> Optional[str]:
    try:
        toks = _tokenize(formula)
    except Exception as e:
        return str(e)
    for idx, (k, v) in enumerate(toks):
        if (k, v) == ("id", "v") and idx + 2 < len(toks) and toks[idx + 2][0] == "str" and toks[idx + 2][1] not in known_keys:
            return f"Unknown variable key '{toks[idx + 2][1]}'"
    _, err = evaluate_formula(formula, lambda k, tp=None: 1)
    return err


# --------------------------------------------------------------------------
# Protocol sanitising
# --------------------------------------------------------------------------
def _s(v, n):
    return (str(v)[:n] if v is not None else "").strip()


def sanitize_variable(v: dict) -> dict:
    opts = []
    for o in (v.get("options") or []):
        if isinstance(o, dict) and o.get("label"):
            opts.append({"code": o.get("code", o["label"]), "label": _s(o["label"], 80)})
        elif isinstance(o, str) and o.strip():
            opts.append({"code": o.strip(), "label": o.strip()})
    out = {
        "key": _s(v.get("key"), 40).lower(),
        "label": _s(v.get("label"), 120) or _s(v.get("key"), 40),
        "description": _s(v.get("description"), 400),
        "level": v.get("level") if v.get("level") in LEVELS else "site",
        "source": v.get("source") if v.get("source") in SOURCES else "study",
        "data_type": v.get("data_type") if v.get("data_type") in DATA_TYPES else "text",
        "unit": _s(v.get("unit"), 20),
        "min": _num(v.get("min")),
        "max": _num(v.get("max")),
        "decimals": int(_num(v.get("decimals")) or 0) if v.get("decimals") not in (None, "") else None,
        "options": opts,
        "required": bool(v.get("required")),
        "case_field_path": _s(v.get("case_field_path"), 80) or None,
        "formula": _s(v.get("formula"), 400) or None,
        "method_note": _s(v.get("method_note"), 400),
        "soft_warning": v.get("soft_warning") if isinstance(v.get("soft_warning"), dict) else None,
    }
    if out["source"] == "linked" and out["case_field_path"] in REGISTRY:
        reg = REGISTRY[out["case_field_path"]]
        out["data_type"] = reg["data_type"]
        out["level"] = reg["level"]
        out["unit"] = reg["unit"] or ""
        if reg["options"]:
            out["options"] = reg["options"]
    return out


def sanitize_timepoint(tp: dict) -> dict:
    return {
        "key": _s(tp.get("key"), 40).lower(),
        "label": _s(tp.get("label"), 80) or _s(tp.get("key"), 40),
        "anchor_event": tp.get("anchor_event") if tp.get("anchor_event") in ANCHORS else "enrolment",
        "anchor_variable_key": _s(tp.get("anchor_variable_key"), 40) or None,
        "offset_days": int(_num(tp.get("offset_days")) or 0),
        "window_before_days": int(_num(tp.get("window_before_days")) or 0),
        "window_after_days": int(_num(tp.get("window_after_days")) or 0),
        "is_baseline": bool(tp.get("is_baseline")),
        "variables": [str(k) for k in (tp.get("variables") or [])],
    }


def sanitize_rule(r: dict) -> dict:
    return {
        "kind": r.get("kind") if r.get("kind") in ("include", "exclude") else "include",
        "mode": r.get("mode") if r.get("mode") in ("query", "checklist") else "query",
        "query": r.get("query") if isinstance(r.get("query"), dict) else None,
        "checklist_text": _s(r.get("checklist_text"), 300),
    }


def validate_protocol(study: dict) -> list[str]:
    """Problems that block activation (BR-07) + structural errors."""
    errs = []
    keys = [v["key"] for v in study.get("variables") or []]
    if len(set(keys)) != len(keys):
        errs.append("Variable keys must be unique")
    for v in study.get("variables") or []:
        if not KEY_RE.match(v["key"]):
            errs.append(f"Variable key '{v['key']}' must be snake_case")
        if v["source"] == "linked" and v["case_field_path"] not in REGISTRY:
            errs.append(f"Variable '{v['key']}': unknown case field '{v['case_field_path']}'")
        if v["source"] == "derived":
            err = validate_formula(v["formula"] or "", set(keys))
            if err:
                errs.append(f"Variable '{v['key']}': formula error — {err}")
        if v["data_type"] in ("single_select", "multi_select") and v["source"] == "study" and not v["options"]:
            errs.append(f"Variable '{v['key']}': select needs options")
    tp_keys = [t["key"] for t in study.get("timepoints") or []]
    if len(set(tp_keys)) != len(tp_keys):
        errs.append("Timepoint keys must be unique")
    for t in study.get("timepoints") or []:
        bad = [k for k in t["variables"] if k not in keys]
        if bad:
            errs.append(f"Timepoint '{t['key']}' references unknown variables: {', '.join(bad)}")
        if t["anchor_event"] == "variable" and t["anchor_variable_key"] not in keys:
            errs.append(f"Timepoint '{t['key']}': anchor variable missing")
    for r in study.get("eligibility") or []:
        if r["mode"] == "query":
            errs += [f"Eligibility rule: {e}" for e in validate_condition(r.get("query"))]
    return errs


def activation_checklist(study: dict) -> list[dict]:
    rules = study.get("eligibility") or []
    tps = study.get("timepoints") or []
    items = [
        ("Research question written", bool((study.get("research_question") or "").strip())),
        ("At least one inclusion rule", any(r["kind"] == "include" for r in rules)),
        ("At least one variable", bool(study.get("variables"))),
        ("Exactly one baseline timepoint", sum(1 for t in tps if t.get("is_baseline")) == 1),
        ("Principal investigator set", bool(study.get("pi_user_id"))),
        ("IEC reference and approval date", bool(study.get("iec_ref")) and bool(study.get("iec_approval_date"))),
        ("Consent form version recorded", bool(study.get("consent_form_version"))),
        ("Protocol has no structural errors", not validate_protocol(study)),
    ]
    if study.get("study_type") == "rct":
        items.append(("CTRI number (RCT)", bool(study.get("ctri_number"))))
    if study.get("is_dissertation"):
        items.append(("Guide approval recorded", bool(study.get("guide_approved_at"))))
    return [{"label": l, "ok": ok} for l, ok in items]


# --------------------------------------------------------------------------
# Schedule engine
# --------------------------------------------------------------------------
def build_forms(study: dict, proc: dict, implant: Optional[dict], enrolled_at: str, existing: Optional[list] = None,
                var_dates: Optional[dict] = None) -> list[dict]:
    """(Re)compute due/window for every timepoint, preserving entered data."""
    prev = {f["timepoint_key"]: f for f in (existing or [])}
    out = []
    for tp in study.get("timepoints") or []:
        f = prev.get(tp["key"]) or {"timepoint_key": tp["key"], "values": {}, "not_done": {}, "status": "not_due",
                                     "visit_date": None, "deviation_flag": False, "deviation_reason": "",
                                     "completed_at": None, "completed_by": None, "locked_at": None, "locked_by": None,
                                     "linked_snapshot": None, "edit_log": []}
        base = anchor_date(proc, implant, tp["anchor_event"], enrolled_at, (var_dates or {}).get(tp.get("anchor_variable_key")))
        if base:
            due = base + timedelta(days=tp["offset_days"])
            f["anchor_date"] = base.isoformat()
            f["due_date"] = due.isoformat()
            f["window_start"] = (due - timedelta(days=tp["window_before_days"])).isoformat()
            f["window_end"] = (due + timedelta(days=tp["window_after_days"])).isoformat()
        else:
            f["anchor_date"] = f["due_date"] = f["window_start"] = f["window_end"] = None
        f["label"] = tp["label"]
        f["is_baseline"] = tp.get("is_baseline", False)
        f["status"] = form_status(f, tp)
        out.append(f)
    return out


def form_status(f: dict, tp: dict) -> str:
    if f.get("locked_at"):
        return "locked"
    if f.get("completed_at"):
        return "completed"
    if f.get("status") == "not_done":
        return "not_done"
    if not f.get("due_date"):
        return "not_due"
    t = today()
    ws, we, due = _to_date(f["window_start"]), _to_date(f["window_end"]), _to_date(f["due_date"])
    if t < ws:
        return "not_due"
    if t <= we:
        return "in_window"
    if (t - due).days > max(1, tp.get("window_after_days", 0)) * MISSED_MULTIPLIER:
        return "missed"
    return "overdue"


# --------------------------------------------------------------------------
# Value validation + derived computation
# --------------------------------------------------------------------------
def validate_value(var: dict, value: Any) -> tuple[Any, Optional[str], Optional[str]]:
    """→ (clean_value, hard_error, soft_warning)."""
    if value in (None, "", []):
        return None, None, None
    dt = var["data_type"]
    if dt in ("integer", "decimal", "scale"):
        n = _num(value)
        if n is None:
            return None, f"{var['label']}: must be a number", None
        if dt == "integer":
            n = int(round(n))
        elif var.get("decimals") is not None:
            n = round(n, var["decimals"])
        if var.get("min") is not None and n < var["min"]:
            return None, f"{var['label']}: below minimum {var['min']}", None
        if var.get("max") is not None and n > var["max"]:
            return None, f"{var['label']}: above maximum {var['max']}", None
        return n, None, soft_warning(var, n)
    if dt == "boolean":
        if isinstance(value, bool):
            return value, None, None
        return str(value).lower() in ("true", "yes", "1"), None, None
    if dt == "single_select":
        codes = {str(o["code"]) for o in var.get("options") or []}
        if codes and str(value) not in codes:
            return None, f"{var['label']}: invalid option", None
        return str(value), None, None
    if dt == "multi_select":
        vals = value if isinstance(value, list) else [value]
        codes = {str(o["code"]) for o in var.get("options") or []}
        vals = [str(x) for x in vals]
        if codes and any(x not in codes for x in vals):
            return None, f"{var['label']}: invalid option", None
        return vals, None, None
    if dt == "date":
        d = _to_date(value)
        if not d:
            return None, f"{var['label']}: invalid date (YYYY-MM-DD)", None
        return d.isoformat(), None, None
    return str(value)[:2000], None, None


def soft_warning(var: dict, n) -> Optional[str]:
    sw = var.get("soft_warning")
    if not sw or not isinstance(n, (int, float)):
        return None
    op, thr, msg = sw.get("op"), _num(sw.get("value")), sw.get("message") or f"{var['label']} outside expected range — confirm"
    if thr is None:
        return None
    hit = {"gt": n > thr, "gte": n >= thr, "lt": n < thr, "lte": n <= thr}.get(op, False)
    return msg if hit else None


def compute_derived(study: dict, forms: list[dict], form: dict, patient_values: dict) -> dict:
    """Return {key: value} for derived variables of `form` (values resolved
    across all forms by timepoint key)."""
    vars_by_key = {v["key"]: v for v in study.get("variables") or []}
    by_tp = {f["timepoint_key"]: f for f in forms}
    cache: dict = {}

    def lookup(key: str, tp_key: Optional[str]):
        target = by_tp.get(tp_key) if tp_key else form
        if target is None:
            return None
        var = vars_by_key.get(key)
        if var is None:
            return None
        if var["level"] == "patient":
            val = (patient_values.get(target["timepoint_key"]) or {}).get(key)
            if val is not None:
                return val
        val = (target.get("values") or {}).get(key)
        if val is None and var["source"] == "derived":
            ck = (key, target["timepoint_key"])
            if ck not in cache:
                cache[ck] = None  # cycle guard
                inner_form = target
                cache[ck], _ = evaluate_formula(var["formula"] or "",
                                                lambda k2, t2=None, _f=inner_form: lookup(k2, t2 or _f["timepoint_key"]))
            return cache[ck]
        return val

    out = {}
    tp_def = next((t for t in study.get("timepoints") or [] if t["key"] == form["timepoint_key"]), None)
    for key in (tp_def or {}).get("variables", []):
        var = vars_by_key.get(key)
        if var and var["source"] == "derived":
            out[key], _ = evaluate_formula(var["formula"] or "", lambda k2, t2=None: lookup(k2, t2))
    return out


# --------------------------------------------------------------------------
# Descriptive statistics
# --------------------------------------------------------------------------
def describe(values: list) -> dict:
    nums = [float(v) for v in values if isinstance(v, (int, float)) and not isinstance(v, bool)]
    if not nums:
        return {"n": 0}
    q = statistics.quantiles(nums, n=4) if len(nums) >= 2 else [nums[0], nums[0], nums[0]]
    return {
        "n": len(nums),
        "mean": round(statistics.fmean(nums), 3),
        "sd": round(statistics.stdev(nums), 3) if len(nums) > 1 else 0.0,
        "median": round(statistics.median(nums), 3),
        "q1": round(q[0], 3), "q3": round(q[2], 3),
        "min": round(min(nums), 3), "max": round(max(nums), 3),
        "ci95": round(1.96 * statistics.stdev(nums) / math.sqrt(len(nums)), 3) if len(nums) > 1 else 0.0,
    }


def count_categories(values: list) -> list[dict]:
    counts: dict = {}
    for v in values:
        for x in (v if isinstance(v, list) else [v]):
            k = "Yes" if x is True else "No" if x is False else str(x)
            counts[k] = counts.get(k, 0) + 1
    total = sum(counts.values()) or 1
    return [{"value": k, "count": c, "pct": round(100 * c / total, 1)} for k, c in sorted(counts.items(), key=lambda kv: -kv[1])]


# --------------------------------------------------------------------------
# CBL seed template — adapted to real registry paths
# --------------------------------------------------------------------------
CBL_TEMPLATE = {
    "id": "cbl_immediate_maxillary_anterior",
    "name": "CBL — Immediate Implants, Maxillary Anterior",
    "description": "Crestal bone loss around immediately placed implants at FDI 13–23 measured on standardised IOPAs at baseline, loading, 6, 12 and 24 months.",
    "code_suggestion": "CBL-IMM-MXA",
    "study_type": "prospective_cohort",
    "primary_unit": "implant",
    "target_n": 30,
    "research_question": "What is the mean crestal bone loss at 12 months around immediately placed implants in the maxillary anterior region?",
    "eligibility": [
        {"kind": "include", "mode": "query", "query": {"all": [
            {"field": "implant.site_fdi", "op": "in", "value": [13, 12, 11, 21, 22, 23]},
            {"field": "case.procedure_type", "op": "eq", "value": "Immediate Implant"},
            {"field": "patient.age_years", "op": "gte", "value": 18}]}},
        {"kind": "include", "mode": "query", "query": {"all": [{"field": "implant.insertion_torque_ncm", "op": "gte", "value": 35}]}},
        {"kind": "include", "mode": "query", "query": {"all": [{"field": "site.socket_type", "op": "in", "value": ["Type I", "Type II"]}]}},
        {"kind": "include", "mode": "checklist", "checklist_text": "Patient willing to attend recall up to 24 months"},
        {"kind": "exclude", "mode": "query", "query": {"any": [
            {"field": "patient.medical.cigarettes_per_day", "op": "gt", "value": 10},
            {"field": "patient.medical.hba1c_pct", "op": "gt", "value": 7},
            {"field": "patient.medical.bisphosphonate_history", "op": "in", "value": ["Oral", "IV"]},
            {"field": "patient.medical.radiation", "op": "eq", "value": "Yes"}]}},
        {"kind": "exclude", "mode": "checklist", "checklist_text": "Active periodontitis at the site"},
    ],
    "variables": [
        {"key": "hba1c", "label": "HbA1c", "level": "patient", "source": "linked", "case_field_path": "patient.medical.hba1c_pct"},
        {"key": "smoker_cigs_day", "label": "Cigarettes per day", "level": "patient", "source": "linked", "case_field_path": "patient.medical.cigarettes_per_day"},
        {"key": "age", "label": "Age", "level": "patient", "source": "linked", "case_field_path": "patient.age_years"},
        {"key": "sex", "label": "Sex", "level": "patient", "source": "linked", "case_field_path": "patient.sex"},
        {"key": "implant_system", "label": "Implant system", "level": "implant", "source": "linked", "case_field_path": "implant.system"},
        {"key": "implant_diameter", "label": "Implant diameter", "level": "implant", "source": "linked", "case_field_path": "implant.diameter_mm"},
        {"key": "implant_length", "label": "Implant length", "level": "implant", "source": "linked", "case_field_path": "implant.length_mm"},
        {"key": "insertion_torque", "label": "Insertion torque", "level": "implant", "source": "linked", "case_field_path": "implant.insertion_torque_ncm"},
        {"key": "bone_type", "label": "Bone type", "level": "implant", "source": "linked", "case_field_path": "implant.bone_type"},
        {"key": "isq", "label": "ISQ", "level": "implant", "source": "study", "data_type": "integer", "unit": "ISQ", "min": 0, "max": 100},
        {"key": "socket_type", "label": "Elian socket type", "level": "site", "source": "linked", "case_field_path": "site.socket_type"},
        {"key": "buccal_plate_mm", "label": "Buccal plate thickness", "level": "site", "source": "study", "data_type": "decimal", "unit": "mm", "min": 0, "max": 5, "decimals": 1,
         "method_note": "CBCT, at crest level, perpendicular to implant axis"},
        {"key": "jump_gap_mm", "label": "Jump gap", "level": "site", "source": "study", "data_type": "decimal", "unit": "mm", "min": 0, "max": 6, "decimals": 1,
         "soft_warning": {"op": "lt", "value": 2, "message": "Gap under 2 mm — confirm grafting decision recorded"}},
        {"key": "gap_graft", "label": "Gap graft material", "level": "site", "source": "study", "data_type": "single_select",
         "options": [{"code": "0", "label": "None"}, {"code": "1", "label": "Xenograft"}, {"code": "2", "label": "Allograft"}, {"code": "3", "label": "Autograft"}, {"code": "4", "label": "Other"}]},
        {"key": "phenotype", "label": "Soft-tissue phenotype", "level": "site", "source": "study", "data_type": "single_select",
         "options": [{"code": "1", "label": "Thin (probe visible)"}, {"code": "2", "label": "Thick"}], "method_note": "Probe transparency at mid-facial"},
        {"key": "bl_mesial", "label": "Bone level mesial", "level": "site", "source": "study", "data_type": "decimal", "unit": "mm", "decimals": 2, "min": -3, "max": 10, "required": True,
         "method_note": "IOPA (paralleling, XCP + bite jig). Shoulder to first bone-to-implant contact, parallel to implant axis; positive = bone apical to shoulder"},
        {"key": "bl_distal", "label": "Bone level distal", "level": "site", "source": "study", "data_type": "decimal", "unit": "mm", "decimals": 2, "min": -3, "max": 10, "required": True,
         "method_note": "Same method as mesial"},
        {"key": "cbl_mesial", "label": "CBL mesial", "level": "site", "source": "derived", "data_type": "decimal", "unit": "mm", "formula": "v('bl_mesial') - v('bl_mesial','baseline')"},
        {"key": "cbl_distal", "label": "CBL distal", "level": "site", "source": "derived", "data_type": "decimal", "unit": "mm", "formula": "v('bl_distal') - v('bl_distal','baseline')"},
        {"key": "cbl_mean", "label": "CBL mean", "level": "site", "source": "derived", "data_type": "decimal", "unit": "mm", "formula": "mean(v('cbl_mesial'), v('cbl_distal'))",
         "soft_warning": {"op": "gt", "value": 1.5, "message": "Exceeds year-1 success threshold (1.5 mm) — confirm measurement"}},
        {"key": "success_flag", "label": "Success (CBL ≤ 1.5 mm at 12 m)", "level": "site", "source": "derived", "data_type": "boolean", "formula": "if(v('cbl_mean','m12') <= 1.5, true, false)",
         "method_note": "Albrektsson-type criterion; threshold editable per study"},
        {"key": "midfacial_change_mm", "label": "Mid-facial mucosal change", "level": "site", "source": "study", "data_type": "decimal", "unit": "mm", "decimals": 1, "min": -5, "max": 5},
        {"key": "pes", "label": "Pink Esthetic Score", "level": "site", "source": "study", "data_type": "scale", "min": 0, "max": 14, "method_note": "Fürhauser 2005"},
        {"key": "pd_mesial", "label": "Probing depth mesial", "level": "site", "source": "study", "data_type": "decimal", "unit": "mm", "min": 0, "max": 15, "decimals": 1},
        {"key": "bop", "label": "Bleeding on probing", "level": "site", "source": "study", "data_type": "boolean"},
        {"key": "complication", "label": "Complication", "level": "implant", "source": "study", "data_type": "single_select",
         "options": [{"code": "none", "label": "None"}, {"code": "mucositis", "label": "Peri-implant mucositis"}, {"code": "periimplantitis", "label": "Peri-implantitis"},
                     {"code": "screw_loosening", "label": "Screw loosening"}, {"code": "fracture", "label": "Prosthesis fracture"}, {"code": "failure", "label": "Implant failure"}]},
    ],
    "timepoints": [
        {"key": "baseline", "label": "Baseline (placement)", "anchor_event": "placement_date", "offset_days": 0, "window_before_days": 0, "window_after_days": 7, "is_baseline": True,
         "variables": ["hba1c", "smoker_cigs_day", "age", "sex", "implant_system", "implant_diameter", "implant_length", "insertion_torque", "bone_type", "isq", "socket_type", "buccal_plate_mm", "jump_gap_mm", "gap_graft", "phenotype", "bl_mesial", "bl_distal"]},
        {"key": "loading", "label": "Loading", "anchor_event": "prosthesis_delivery_date", "offset_days": 0, "window_before_days": 0, "window_after_days": 14,
         "variables": ["isq", "bl_mesial", "bl_distal", "cbl_mesial", "cbl_distal", "cbl_mean", "midfacial_change_mm", "pes", "complication"]},
        {"key": "m6", "label": "6 months", "anchor_event": "prosthesis_delivery_date", "offset_days": 182, "window_before_days": 15, "window_after_days": 15,
         "variables": ["bl_mesial", "bl_distal", "cbl_mesial", "cbl_distal", "cbl_mean", "pd_mesial", "bop", "complication"]},
        {"key": "m12", "label": "12 months", "anchor_event": "prosthesis_delivery_date", "offset_days": 365, "window_before_days": 30, "window_after_days": 30,
         "variables": ["bl_mesial", "bl_distal", "cbl_mesial", "cbl_distal", "cbl_mean", "success_flag", "midfacial_change_mm", "pes", "pd_mesial", "bop", "complication"]},
        {"key": "m24", "label": "24 months", "anchor_event": "prosthesis_delivery_date", "offset_days": 730, "window_before_days": 60, "window_after_days": 60,
         "variables": ["bl_mesial", "bl_distal", "cbl_mesial", "cbl_distal", "cbl_mean", "midfacial_change_mm", "pes", "pd_mesial", "bop", "complication"]},
    ],
}

ISQ_TEMPLATE = {
    "id": "isq_stability_over_time",
    "name": "ISQ Stability Over Time",
    "description": "Implant stability quotient at placement, second stage and loading; relation to bone type and torque.",
    "code_suggestion": "ISQ-STAB",
    "study_type": "prospective_cohort", "primary_unit": "implant", "target_n": 40,
    "research_question": "How does implant stability (ISQ) change from placement to loading, and does bone type predict the change?",
    "eligibility": [
        {"kind": "include", "mode": "query", "query": {"all": [{"field": "implant.site_fdi", "op": "exists"}, {"field": "patient.age_years", "op": "gte", "value": 18}]}},
        {"kind": "include", "mode": "checklist", "checklist_text": "Osstell / RFA device available for all visits"},
    ],
    "variables": [
        {"key": "age", "label": "Age", "level": "patient", "source": "linked", "case_field_path": "patient.age_years"},
        {"key": "diabetes", "label": "Diabetes", "level": "patient", "source": "linked", "case_field_path": "patient.medical.diabetes"},
        {"key": "bone_type", "label": "Bone type", "level": "implant", "source": "linked", "case_field_path": "implant.bone_type"},
        {"key": "insertion_torque", "label": "Insertion torque", "level": "implant", "source": "linked", "case_field_path": "implant.insertion_torque_ncm"},
        {"key": "implant_diameter", "label": "Diameter", "level": "implant", "source": "linked", "case_field_path": "implant.diameter_mm"},
        {"key": "implant_length", "label": "Length", "level": "implant", "source": "linked", "case_field_path": "implant.length_mm"},
        {"key": "isq_bl", "label": "ISQ bucco-lingual", "level": "implant", "source": "study", "data_type": "integer", "unit": "ISQ", "min": 0, "max": 100, "required": True},
        {"key": "isq_md", "label": "ISQ mesio-distal", "level": "implant", "source": "study", "data_type": "integer", "unit": "ISQ", "min": 0, "max": 100, "required": True},
        {"key": "isq_mean", "label": "ISQ mean", "level": "implant", "source": "derived", "data_type": "decimal", "unit": "ISQ", "formula": "mean(v('isq_bl'), v('isq_md'))"},
        {"key": "isq_change", "label": "ISQ change from placement", "level": "implant", "source": "derived", "data_type": "decimal", "unit": "ISQ", "formula": "v('isq_mean') - v('isq_mean','placement')"},
    ],
    "timepoints": [
        {"key": "placement", "label": "Placement", "anchor_event": "placement_date", "offset_days": 0, "window_before_days": 0, "window_after_days": 3, "is_baseline": True,
         "variables": ["age", "diabetes", "bone_type", "insertion_torque", "implant_diameter", "implant_length", "isq_bl", "isq_md", "isq_mean"]},
        {"key": "second_stage", "label": "Second stage", "anchor_event": "second_stage_date", "offset_days": 0, "window_before_days": 7, "window_after_days": 7,
         "variables": ["isq_bl", "isq_md", "isq_mean", "isq_change"]},
        {"key": "loading", "label": "Loading", "anchor_event": "prosthesis_delivery_date", "offset_days": 0, "window_before_days": 7, "window_after_days": 14,
         "variables": ["isq_bl", "isq_md", "isq_mean", "isq_change"]},
    ],
}

PES_WES_TEMPLATE = {
    "id": "pes_wes_aesthetic_outcome",
    "name": "PES / WES Aesthetic Outcome",
    "description": "Pink and White Esthetic Scores of single anterior implant crowns at delivery and 12 months.",
    "code_suggestion": "PES-WES",
    "study_type": "prospective_cohort", "primary_unit": "site", "target_n": 30,
    "research_question": "What is the aesthetic outcome (PES/WES) of single implant crowns in the anterior maxilla at 12 months?",
    "eligibility": [
        {"kind": "include", "mode": "query", "query": {"all": [{"field": "implant.site_fdi", "op": "in", "value": [13, 12, 11, 21, 22, 23]}]}},
        {"kind": "include", "mode": "checklist", "checklist_text": "Standardised frontal photograph available at each visit"},
    ],
    "variables": [
        {"key": "era", "label": "Aesthetic risk (ERA)", "level": "patient", "source": "linked", "case_field_path": "case.aesthetic_risk.overall"},
        {"key": "biotype", "label": "Gingival biotype", "level": "patient", "source": "linked", "case_field_path": "case.gingival_biotype"},
        {"key": "smile_line", "label": "Smile line", "level": "patient", "source": "linked", "case_field_path": "case.smile_line"},
        {"key": "pes", "label": "Pink Esthetic Score", "level": "site", "source": "study", "data_type": "scale", "min": 0, "max": 14, "required": True, "method_note": "Fürhauser 2005 — 7 items × 0–2"},
        {"key": "wes", "label": "White Esthetic Score", "level": "site", "source": "study", "data_type": "scale", "min": 0, "max": 10, "required": True, "method_note": "Belser 2009 — 5 items × 0–2"},
        {"key": "pes_wes_total", "label": "PES + WES", "level": "site", "source": "derived", "data_type": "integer", "formula": "v('pes') + v('wes')"},
        {"key": "aesthetic_success", "label": "Aesthetic success (PES ≥ 8 and WES ≥ 6)", "level": "site", "source": "derived", "data_type": "boolean", "formula": "if(v('pes') >= 8 and v('wes') >= 6, true, false)"},
        {"key": "patient_satisfaction", "label": "Patient satisfaction (VAS 0–10)", "level": "patient", "source": "study", "data_type": "scale", "min": 0, "max": 10},
    ],
    "timepoints": [
        {"key": "delivery", "label": "Crown delivery", "anchor_event": "prosthesis_delivery_date", "offset_days": 0, "window_before_days": 0, "window_after_days": 14, "is_baseline": True,
         "variables": ["era", "biotype", "smile_line", "pes", "wes", "pes_wes_total", "aesthetic_success", "patient_satisfaction"]},
        {"key": "m12", "label": "12 months", "anchor_event": "prosthesis_delivery_date", "offset_days": 365, "window_before_days": 30, "window_after_days": 30,
         "variables": ["pes", "wes", "pes_wes_total", "aesthetic_success", "patient_satisfaction"]},
    ],
}

TEMPLATES = [CBL_TEMPLATE, ISQ_TEMPLATE, PES_WES_TEMPLATE]


# --------------------------------------------------------------------------
# Synopsis / dissertation document (D3). Section order follows the common
# Indian university synopsis format; free-text sections are editable per study.
# --------------------------------------------------------------------------
SYNOPSIS_FREE_SECTIONS = [
    ("introduction", "Introduction"),
    ("review_of_literature", "Review of Literature"),
    ("aims_objectives", "Aims and Objectives"),
    ("materials_methods_extra", "Materials and Methods — additional details"),
    ("statistical_analysis", "Statistical Analysis"),
    ("ethical_considerations", "Ethical Considerations"),
    ("references", "References"),
]


def sanitize_synopsis(body: dict) -> dict:
    body = body or {}
    out = {k: _s(body.get(k), 12000) for k, _ in SYNOPSIS_FREE_SECTIONS}
    for k, n in (("university_name", 160), ("college_name", 160), ("department_name", 120), ("degree", 80), ("academic_year", 20),
                 ("student_name", 100), ("guide_name", 100), ("co_guide_name", 100), ("hod_name", 100), ("sample_size_text", 2000)):
        out[k] = _s(body.get(k), n)
    return out


def _eligibility_lines(st: dict, kind: str) -> list[str]:
    return [describe_condition(r["query"]) if r["mode"] == "query" else r.get("checklist_text", "")
            for r in st.get("eligibility", []) if r["kind"] == kind]


def _methods_paragraphs(st: dict) -> list[str]:
    tps = st.get("timepoints", [])
    unit = st.get("primary_unit", "implant")
    outcome_vars = [v for v in st.get("variables", []) if v["source"] != "linked"]
    linked = [v for v in st.get("variables", []) if v["source"] == "linked"]
    design = st.get("study_type", "").replace("_", " ")
    p = [
        f"Study design: This is a {design} with the {unit} as the primary unit of analysis. "
        f"The study will be conducted in the Department of Implantology using the Implanr clinical record system; "
        f"a target of {st.get('target_n')} {unit}s will be enrolled.",
        "Participants will be screened against the inclusion and exclusion criteria listed above. Eligibility against coded "
        "case fields is evaluated automatically; items requiring clinical confirmation are verified at enrolment. Written "
        f"informed consent (form version {st.get('consent_form_version') or '—'}) will be obtained in the participant's language.",
    ]
    if tps:
        sched = "; ".join(f"{t['label']} ({t['anchor_event'].replace('_', ' ')}{' + ' + str(t['offset_days']) + ' days' if t['offset_days'] else ''}, "
                          f"window −{t['window_before_days']}/+{t['window_after_days']} days)" for t in tps)
        p.append(f"Assessment schedule: Data will be collected at {len(tps)} timepoints — {sched}. Visits outside the protocol window are recorded as deviations.")
    if linked:
        p.append("Baseline characteristics extracted from the clinical record: " + ", ".join(v["label"] for v in linked) + ".")
    if outcome_vars:
        lines = []
        for v in outcome_vars:
            s = v["label"] + (f" ({v['unit']})" if v.get("unit") else "")
            if v["source"] == "derived":
                s += f" — derived: {v.get('formula')}"
            if v.get("method_note"):
                s += f". Method: {v['method_note']}"
            lines.append(s)
        p.append("Study variables and measurement methods: " + " | ".join(lines) + ".")
    p.append("Data management: Data are captured in structured electronic forms with range checks, mandatory-field rules and an "
             "immutable audit trail. Linked clinical values are read-only and snapshotted at form lock. Exports are de-identified "
             "(subject codes; dates expressed as days from the anchor event).")
    return p


def _default_stats(st: dict) -> str:
    return ("Descriptive statistics (mean ± SD, median [IQR] for continuous variables; counts and percentages for categorical variables) "
            "will be computed per timepoint. Change from baseline will be analysed with paired tests or repeated-measures models as "
            "appropriate. Where multiple implants are clustered within patients, mixed-effects models / GEE will be used. "
            "Statistical significance will be set at p < 0.05.")


def _synopsis_sections(st: dict, dash: Optional[dict]) -> list[tuple[str, list[str]]]:
    syn = st.get("synopsis") or {}
    inc, exc = _eligibility_lines(st, "include"), _eligibility_lines(st, "exclude")
    sections: list[tuple[str, list[str]]] = [("Title", [st.get("title", "")])]
    investigators = [f"Principal investigator: {st.get('pi_name', '')}"]
    if st.get("is_dissertation"):
        investigators += [f"Candidate: {syn.get('student_name') or st.get('pi_name', '')}", f"Guide: {syn.get('guide_name') or st.get('guide_name') or '—'}"]
        if syn.get("co_guide_name"):
            investigators.append(f"Co-guide: {syn['co_guide_name']}")
    sections.append(("Investigators", investigators))
    sections.append(("Introduction", [syn.get("introduction") or "(to be written)"]))
    if syn.get("review_of_literature"):
        sections.append(("Review of Literature", [syn["review_of_literature"]]))
    sections.append(("Research Question", [st.get("research_question", "")]))
    sections.append(("Aims and Objectives", [syn.get("aims_objectives") or f"To determine: {st.get('research_question', '')}"]))
    sections.append(("Materials and Methods", _methods_paragraphs(st) + ([syn["materials_methods_extra"]] if syn.get("materials_methods_extra") else [])))
    sections.append(("Inclusion Criteria", [f"{i + 1}. {x}" for i, x in enumerate(inc)] or ["—"]))
    sections.append(("Exclusion Criteria", [f"{i + 1}. {x}" for i, x in enumerate(exc)] or ["—"]))
    sections.append(("Sample Size", [syn.get("sample_size_text") or f"A total of {st.get('target_n')} {st.get('primary_unit')}s will be enrolled."]))
    sections.append(("Statistical Analysis", [syn.get("statistical_analysis") or _default_stats(st)]))
    ethics = [f"IEC reference: {st.get('iec_ref') or '—'}; approval date: {st.get('iec_approval_date') or '—'}."]
    if st.get("ctri_number"):
        ethics.append(f"CTRI registration: {st['ctri_number']}.")
    ethics.append(syn.get("ethical_considerations") or "Written informed consent will be obtained from every participant; participants may withdraw at any time without affecting their clinical care. Data are held de-identified and access is audit-logged.")
    sections.append(("Ethical Considerations", ethics))
    if dash:
        flow = dash.get("flow", {})
        sections.append(("Current Status", [f"Enrolled {flow.get('enrolled', 0)} of {st.get('target_n')} ({flow.get('patients', 0)} patients); completed forms: {dash.get('completed_forms', 0)}."]))
    sections.append(("References", [syn.get("references") or "(to be added)"]))
    return sections


def _crf_rows(st: dict) -> list[list[str]]:
    tps = st.get("timepoints", [])
    rows = [["Variable", "Level", "Source", "Type / unit"] + [t["label"] for t in tps]]
    for v in st.get("variables", []):
        rows.append([v["label"], v["level"], v["source"], (v["data_type"] + (f" ({v['unit']})" if v.get("unit") else ""))] +
                    ["✓" if v["key"] in t["variables"] else "" for t in tps])
    return rows


def _title_block(st: dict) -> list[str]:
    syn = st.get("synopsis") or {}
    lines = [x for x in (syn.get("university_name"), syn.get("college_name"), syn.get("department_name")) if x]
    if syn.get("degree"):
        lines.append(f"Synopsis for {syn['degree']}" + (f" — {syn['academic_year']}" if syn.get("academic_year") else ""))
    return lines


def build_synopsis_docx(st: dict, dash: Optional[dict] = None) -> bytes:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    style = doc.styles["Normal"]; style.font.name = "Times New Roman"; style.font.size = Pt(12)
    for line in _title_block(st):
        p = doc.add_paragraph(); r = p.add_run(line.upper()); r.bold = True; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(); r = p.add_run(st.get("title", "")); r.bold = True; r.font.size = Pt(16); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(f"Study code: {st.get('code')} · Protocol v{st.get('protocol_version', '1.0')} · Generated {today().isoformat()}"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for heading, paras in _synopsis_sections(st, dash):
        doc.add_heading(heading, level=1)
        for text in paras:
            for chunk in str(text).split("\n"):
                if chunk.strip():
                    doc.add_paragraph(chunk.strip())
    doc.add_page_break()
    doc.add_heading("Annexure A — Case Record Form (variables × timepoints)", level=1)
    rows = _crf_rows(st)
    table = doc.add_table(rows=len(rows), cols=len(rows[0])); table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = table.cell(i, j); c.text = str(cell)
            for par in c.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(8); run.bold = i == 0
    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()


def build_synopsis_pdf(st: dict, dash: Optional[dict] = None) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    from xml.sax.saxutils import escape
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=25 * mm, rightMargin=25 * mm, topMargin=20 * mm, bottomMargin=20 * mm)
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Times-Roman", fontSize=11, leading=15, spaceAfter=6)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Times-Bold", fontSize=13, spaceBefore=10, spaceAfter=4)
    center = ParagraphStyle("c", parent=body, alignment=1, fontName="Times-Bold")
    story = [Paragraph(escape(x.upper()), center) for x in _title_block(st)]
    story += [Spacer(1, 6), Paragraph(f"<font size=15>{escape(st.get('title', ''))}</font>", center),
              Paragraph(escape(f"Study code: {st.get('code')} · Protocol v{st.get('protocol_version', '1.0')} · Generated {today().isoformat()}"), ParagraphStyle("s", parent=body, alignment=1, fontSize=9))]
    for heading, paras in _synopsis_sections(st, dash):
        story.append(Paragraph(escape(heading), h1))
        for text in paras:
            for chunk in str(text).split("\n"):
                if chunk.strip():
                    story.append(Paragraph(escape(chunk.strip()), body))
    story.append(PageBreak()); story.append(Paragraph("Annexure A — Case Record Form (variables × timepoints)", h1))
    rows = [[Paragraph(escape(str(c)), ParagraphStyle("t", parent=body, fontSize=7, leading=9)) for c in r] for r in _crf_rows(st)]
    t = Table(rows, repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E3F2FD")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(t)
    doc.build(story)
    return buf.getvalue()
